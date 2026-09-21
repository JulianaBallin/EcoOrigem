"""Small helpers on top of python-docx for the technical report."""

# pylint: disable=too-many-arguments,too-many-positional-arguments,too-many-locals

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

FONT = "Arial"
GREEN_950, GREEN_900, GREEN_800, GREEN_700 = "062A19", "0B3D24", "0F5132", "17693D"
GREEN_100, GREEN_50, BORDER, MUTED = "E3F0DC", "F1F7EE", "B9CDB4", "5A6B61"
TEXT_WIDTH_CM = 16.0


def rgb(hex_color: str) -> RGBColor:
    """Convert a hex string to an RGBColor."""
    return RGBColor.from_string(hex_color)


def set_font(
    run_or_style, size: float | None = None, bold=None, italic=None, color=None
):
    """Apply the project font to a run or a style."""
    font = run_or_style.font
    font.name = FONT
    element = run_or_style.element if hasattr(run_or_style, "element") else None
    if element is not None:
        rpr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else None
        if rpr is not None:
            fonts = rpr.get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(attr), FONT)
            for attr in (
                "w:asciiTheme",
                "w:hAnsiTheme",
                "w:eastAsiaTheme",
                "w:cstheme",
            ):
                if fonts.get(qn(attr)) is not None:
                    del fonts.attrib[qn(attr)]
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color is not None:
        font.color.rgb = rgb(color)


def shade(cell, fill: str) -> None:
    """Fill a table cell with a background color."""
    props = cell._tc.get_or_add_tcPr()  # pylint: disable=protected-access
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def cell_borders(
    cell, color: str = BORDER, size: int = 4, sides=("top", "left", "bottom", "right")
) -> None:
    """Set single-line borders on a cell."""
    props = cell._tc.get_or_add_tcPr()  # pylint: disable=protected-access
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        if side in sides:
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), str(size))
            border.set(qn("w:color"), color)
        else:
            border.set(qn("w:val"), "nil")
        borders.append(border)
    props.append(borders)


def cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    """Set inner cell margins in twentieths of a point."""
    props = cell._tc.get_or_add_tcPr()  # pylint: disable=protected-access
    margins = OxmlElement("w:tcMar")
    for name, value in (
        ("top", top),
        ("left", left),
        ("bottom", bottom),
        ("right", right),
    ):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)


def fixed_table(table, widths_cm: list[float]) -> None:
    """Force a fixed layout with explicit column widths."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)  # pylint: disable=protected-access
    for row in table.rows:
        for index, width in enumerate(widths_cm):
            row.cells[index].width = Cm(width)
    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)


def repeat_header(row) -> None:
    """Repeat a table row at the top of every page."""
    props = row._tr.get_or_add_trPr()  # pylint: disable=protected-access
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def no_split(row) -> None:
    """Keep a table row on a single page."""
    props = row._tr.get_or_add_trPr()  # pylint: disable=protected-access
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    props.append(node)


def bottom_border(
    paragraph, color: str = GREEN_700, size: int = 8, space: int = 4
) -> None:
    """Draw a line under a paragraph."""
    ppr = paragraph._p.get_or_add_pPr()  # pylint: disable=protected-access
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def top_border(paragraph, color: str = BORDER, size: int = 6, space: int = 4) -> None:
    """Draw a line above a paragraph."""
    ppr = paragraph._p.get_or_add_pPr()  # pylint: disable=protected-access
    borders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), str(space))
    top.set(qn("w:color"), color)
    borders.append(top)
    ppr.append(borders)


def add_field(
    paragraph, instruction: str, size: float = 8.5, color: str = MUTED
) -> None:
    """Insert a simple field (PAGE, NUMPAGES) into a paragraph."""
    run = paragraph.add_run()
    set_font(run, size, color=color)
    for kind, text in (
        ("begin", None),
        (None, instruction),
        ("separate", None),
        (None, "1"),
        ("end", None),
    ):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            run._r.append(node)  # pylint: disable=protected-access
        elif text == instruction:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = f" {instruction} "
            run._r.append(node)  # pylint: disable=protected-access
        else:
            node = OxmlElement("w:t")
            node.text = text
            run._r.append(node)  # pylint: disable=protected-access


class Report:
    """Fluent wrapper that owns the document and its numbering counters."""

    def __init__(self) -> None:
        self.doc = Document()
        self.figure_number = 0
        self.table_number = 0
        self.headings: list[tuple[int, str]] = []
        self._configure()

    # ------------------------------------------------------------ setup
    def _configure(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.orientation = WD_ORIENT.PORTRAIT
        section.top_margin = section.bottom_margin = Cm(0)
        section.left_margin = section.right_margin = Cm(0)
        section.header_distance = section.footer_distance = Cm(0)

        normal = self.doc.styles["Normal"]
        set_font(normal, 11)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.line_spacing = 1.3
        normal.paragraph_format.space_after = Pt(6)
        for name, size, color, before, after in (
            ("Heading 1", 16, GREEN_900, 18, 8),
            ("Heading 2", 13, GREEN_800, 14, 5),
            ("Heading 3", 11.5, GREEN_700, 10, 4),
        ):
            style = self.doc.styles[name]
            set_font(style, size, bold=True, italic=False, color=color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.line_spacing = 1.1
        for name in ("List Bullet", "List Number"):
            style = self.doc.styles[name]
            set_font(style, 11)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            style.paragraph_format.space_after = Pt(3)
            style.paragraph_format.line_spacing = 1.25

    def start_body(self) -> None:
        """Close the full-bleed cover section and open the body section."""
        self.doc.add_section(WD_SECTION.NEW_PAGE)
        marker = self.doc.paragraphs[-1]
        marker.paragraph_format.space_after = Pt(0)
        marker.paragraph_format.space_before = Pt(0)
        marker.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        marker.paragraph_format.line_spacing = Pt(1)
        section = self.doc.sections[-1]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(3), Cm(2)
        section.left_margin, section.right_margin = Cm(3), Cm(2)
        section.header_distance, section.footer_distance = Cm(1.25), Cm(1.0)
        self._header_footer(section)

    @staticmethod
    def _right_tab(paragraph) -> None:
        """Clear the default header tabs and right-align the text width."""
        stops = paragraph.paragraph_format.tab_stops
        for position in (4680, 9360):
            stops.add_tab_stop(Twips(position), WD_TAB_ALIGNMENT.CLEAR)
        stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)

    def _header_footer(self, section) -> None:
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        self._right_tab(paragraph)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run("EcoOrigem | Relatório técnico")
        set_font(run, 8.5, bold=True, color=GREEN_800)
        run = paragraph.add_run("\tOficina de Desenvolvimento de Sistemas III")
        set_font(run, 8.5, color=MUTED)
        bottom_border(paragraph, GREEN_700, 8, 4)
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        self._right_tab(paragraph)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        top_border(paragraph)
        run = paragraph.add_run("Universidade do Estado do Amazonas\t")
        set_font(run, 8.5, color=MUTED)
        run = paragraph.add_run("Página ")
        set_font(run, 8.5, color=MUTED)
        add_field(paragraph, "PAGE")
        run = paragraph.add_run(" de ")
        set_font(run, 8.5, color=MUTED)
        add_field(paragraph, "NUMPAGES")

    # ---------------------------------------------------------- content
    def heading(self, level: int, text: str, page_break: bool = False):
        """Add a heading and remember it for the table of contents."""
        paragraph = self.doc.add_heading(text, level=level)
        paragraph.paragraph_format.page_break_before = page_break
        if level == 1:
            bottom_border(paragraph, GREEN_100, 10, 3)
        self.headings.append((level, text))
        return paragraph

    def paragraph(
        self,
        text: str = "",
        bold_lead: str | None = None,
        align=None,
        size: float | None = None,
        italic=False,
        color=None,
        space_after=None,
        keep_next=False,
    ):
        """Add a justified body paragraph, optionally with a bold lead-in."""
        paragraph = self.doc.add_paragraph()
        if bold_lead:
            run = paragraph.add_run(bold_lead + " ")
            set_font(run, size, bold=True, color=color)
        run = paragraph.add_run(text)
        set_font(run, size, italic=italic, color=color)
        if align is not None:
            paragraph.paragraph_format.alignment = align
        if space_after is not None:
            paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.keep_with_next = keep_next
        return paragraph

    def bullets(
        self, items: list[str | tuple[str, str]], numbered: bool = False
    ) -> None:
        """Add a bulleted or numbered list. Tuples are (bold lead, text)."""
        style = "List Number" if numbered else "List Bullet"
        for item in items:
            paragraph = self.doc.add_paragraph(style=style)
            if isinstance(item, tuple):
                run = paragraph.add_run(item[0] + " ")
                set_font(run, bold=True)
                run = paragraph.add_run(item[1])
                set_font(run)
            else:
                run = paragraph.add_run(item)
                set_font(run)

    def code(self, lines: list[str]) -> None:
        """Add a block of commands in a shaded, single-cell table."""
        table = self.doc.add_table(rows=1, cols=1)
        fixed_table(table, [TEXT_WIDTH_CM])
        cell = table.rows[0].cells[0]
        shade(cell, "0F1F17")
        cell_margins(cell, 100, 100, 160, 160)
        cell.paragraphs[0].text = ""
        for index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(line)
            run.font.name = "Liberation Mono"
            run.font.size = Pt(8.5)
            run.font.color.rgb = rgb("D7F0C8")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(
        self, path: Path, title: str, caption: str, width_cm: float = 15.0
    ) -> None:
        """Add a figure with a centered title above and an explanatory caption below."""
        self.figure_number += 1
        heading = self.doc.add_paragraph()
        heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(4)
        run = heading.add_run(f"Figura {self.figure_number} - {title}")
        set_font(run, 10, bold=True, color=GREEN_900)
        picture = self.doc.add_paragraph()
        picture.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        picture.paragraph_format.space_after = Pt(4)
        picture.add_run().add_picture(str(path), width=Cm(width_cm))
        legend = self.doc.add_paragraph()
        legend.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        legend.paragraph_format.space_after = Pt(12)
        legend.paragraph_format.left_indent = Cm(0.8)
        legend.paragraph_format.right_indent = Cm(0.8)
        run = legend.add_run("Legenda: " + caption)
        set_font(run, 9, italic=True, color=MUTED)

    def table(
        self,
        title: str,
        headers: list[str],
        rows: list[list[str]],
        widths_cm: list[float],
        caption: str | None = None,
        size: float = 9,
        align_center: set[int] | None = None,
        keep_together: bool = True,
    ) -> None:
        """Add a titled table with a shaded header row and a legend."""
        self.table_number += 1
        heading = self.doc.add_paragraph()
        heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(4)
        run = heading.add_run(f"Tabela {self.table_number} - {title}")
        set_font(run, 10, bold=True, color=GREEN_900)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        fixed_table(table, widths_cm)
        align_center = align_center or set()
        for index, text in enumerate(headers):
            cell = table.rows[0].cells[index]
            shade(cell, GREEN_800)
            cell_borders(cell, GREEN_800)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index in align_center
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            run = paragraph.add_run(text)
            set_font(run, size, bold=True, color="FFFFFF")
        repeat_header(table.rows[0])
        for r_index, row in enumerate(rows, start=1):
            no_split(table.rows[r_index])
            for c_index, text in enumerate(row):
                cell = table.rows[r_index].cells[c_index]
                if r_index % 2 == 0:
                    shade(cell, GREEN_50)
                cell_borders(cell)
                cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if c_index in align_center
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                bold = text.startswith("**")
                run = paragraph.add_run(text.strip("*") if bold else text)
                set_font(run, size, bold=bold)
        if keep_together:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for cell_paragraph in cell.paragraphs:
                        cell_paragraph.paragraph_format.keep_with_next = True
        legend = self.doc.add_paragraph()
        legend.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        legend.paragraph_format.space_before = Pt(4)
        legend.paragraph_format.space_after = Pt(12)
        if caption:
            run = legend.add_run("Legenda: " + caption)
            set_font(run, 9, italic=True, color=MUTED)

    def page_break(self) -> None:
        """Insert a page break."""
        self.doc.add_page_break()

    def toc_entry(self, level: int, text: str, page: str) -> None:
        """Add a table of contents line with dot leader and page number."""
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Cm(0.7 * (level - 1))
        paragraph.paragraph_format.space_after = Pt(3 if level > 1 else 4)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        run = paragraph.add_run(text)
        set_font(run, 10.5 if level == 1 else 10, bold=level == 1)
        run = paragraph.add_run(f"\t{page}")
        set_font(run, 10.5 if level == 1 else 10, bold=level == 1)

    def save(self, path: Path) -> None:
        """Write the document."""
        self.doc.save(str(path))
