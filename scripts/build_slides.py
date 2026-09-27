"""Build the presentation (PPTX and PDF) used in the class demonstration.

Flow: context, problem, data, solution, architecture, contract, demonstration,
tests and conclusion, followed by bonus slides about what was delivered beyond
the assignment. Numbers come from ``docs/evidencias/metricas.json`` (written by
``make report``) and official statistics are cited on the slide.
Run with ``make slides``.
"""

# pylint: disable=line-too-long,too-many-locals,too-many-statements,too-many-arguments,too-many-positional-arguments,protected-access,no-member,too-many-lines,too-many-branches,unused-argument

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EVID = DOCS / "assets" / "evidencias"
DIAG = DOCS / "assets" / "diagramas"
PRODUCTS = DOCS / "assets" / "produtos"
SLIDE_IMG = DOCS / "assets" / "slides"
LOGO = ROOT / "assets" / "img" / "logo.png"
LOGO_SOURCE = ROOT / "assets" / "img" / "ecoorigem-logo.png"
UEA = DOCS / "assets" / "logos" / "uea.png"
METRICS = DOCS / "evidencias" / "metricas.json"
REPORT_PDF = DOCS / "relatorio" / "relatorio-tecnico-ecoorigem.pdf"
OUT_DIR = DOCS / "slides"
OUT = OUT_DIR / "apresentacao-ecoorigem.pptx"

FONT, MONO = "Arial", "Courier New"
INK, FOREST, GREEN, LIME = "17261D", "0B3D24", "17693D", "8CC63F"
TINT, MUTED, LINE, RED, PALE = "EEF5EA", "5A6B61", "D2DECE", "B3261E", "CFE5C5"
COURSE = "Oficina de Desenvolvimento de Sistemas III"
DPI = 200


def rgb(value: str) -> RGBColor:
    """Hex string to RGBColor."""
    return RGBColor.from_string(value)


def font_file(bold: bool = False) -> str:
    """Path of a Liberation Sans font file (falls back to DejaVu)."""
    query = "Liberation Sans:bold" if bold else "Liberation Sans"
    result = subprocess.run(
        ["fc-match", "-f", "%{file}", query],
        capture_output=True,
        text=True,
        check=False,
    )
    return result.stdout or "DejaVuSans.ttf"


# ------------------------------------------------------------ image assets
def gradient(
    size: tuple[int, int], start: str, end: str, weight: float = 0.6
) -> Image.Image:
    """Soft diagonal gradient with light dithering to avoid banding."""
    w, h = size
    x = np.linspace(0, 1, w)[None, :]
    y = np.linspace(0, 1, h)[:, None]
    t = np.clip((1 - weight) * x + weight * y, 0, 1)[..., None]
    a = np.array([int(start[i : i + 2], 16) for i in (0, 2, 4)], dtype=float)
    b = np.array([int(end[i : i + 2], 16) for i in (0, 2, 4)], dtype=float)
    data = a * (1 - t) + b * t
    data += np.random.default_rng(7).uniform(-0.9, 0.9, data.shape)
    return Image.fromarray(np.clip(data, 0, 255).astype("uint8"), "RGB")


def keyed_logo(
    box: tuple[int, int, int, int], drop_below: int | None = None, cutoff: int = 236
) -> Image.Image:
    """Crop the project logo from the artwork and make the pale background transparent.

    ``drop_below`` removes forest fragments that enter the crop right of the emblem.
    """
    crop = Image.open(LOGO_SOURCE).convert("RGB").crop(box)
    arr = np.asarray(crop).astype(float)
    alpha = np.clip((cutoff - arr.min(axis=2)) / 55.0, 0, 1)
    if drop_below is not None:
        alpha[drop_below:, 520:] = 0
    return Image.fromarray(np.dstack([arr, alpha * 255]).astype("uint8"), "RGBA")


def forest_watermark(canvas: Image.Image, opacity: float = 0.2) -> Image.Image:
    """Paste the forest silhouette of the artwork, very softly, along the bottom."""
    source = Image.open(LOGO_SOURCE).convert("RGB")
    top = 605
    forest = source.crop((0, top, source.width, source.height))
    arr = np.asarray(forest).astype(float)
    dark = np.clip((205 - arr.min(axis=2)) / (205 - 40), 0, 1) ** 1.1
    # the emblem of the artwork reaches into this band: fade it out
    emblem_h, emblem_w = 712 - top, 640
    fade = np.ones_like(dark)
    fade[:emblem_h, :emblem_w] = 0
    fade[emblem_h : emblem_h + 30, :emblem_w] = np.linspace(0, 1, 30)[:, None]
    dark = dark * fade
    scale = canvas.width / forest.width
    mask = Image.fromarray((dark * 255 * opacity).astype("uint8"), "L").resize(
        (canvas.width, int(forest.height * scale)), Image.LANCZOS
    )
    tint = Image.new("RGB", mask.size, "#" + FOREST)
    canvas = canvas.copy()
    canvas.paste(tint, (0, canvas.height - mask.height), mask)
    return canvas


def cover_crop(path: Path, w_in: float, h_in: float, focus=(0.5, 0.5)) -> Image.Image:
    """Crop an image to fill a box of the given size in inches."""
    image = Image.open(path).convert("RGB")
    target = w_in / h_in
    if image.width / image.height > target:
        new_w = int(image.height * target)
        left = int((image.width - new_w) * focus[0])
        image = image.crop((left, 0, left + new_w, image.height))
    else:
        new_h = int(image.width / target)
        top = int((image.height - new_h) * focus[1])
        image = image.crop((0, top, image.width, top + new_h))
    return image.resize((int(w_in * DPI), int(h_in * DPI)), Image.LANCZOS)


def rounded(image: Image.Image, radius_in: float) -> Image.Image:
    """Apply rounded corners with an alpha channel."""
    mask = Image.new("L", image.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle(
        (0, 0, image.width - 1, image.height - 1), int(radius_in * DPI), fill=255
    )
    out = image.convert("RGBA")
    out.putalpha(mask)
    return out


def caption_band(image: Image.Image, text: str) -> Image.Image:
    """Burn a soft dark caption band with text into the lower edge of a photo."""
    image = image.convert("RGBA")
    band_h = int(0.62 * DPI)
    overlay = Image.new("RGBA", image.size, (0, 0, 0, 0))
    grad = np.zeros((band_h, image.width, 4), dtype="uint8")
    ramp = np.linspace(0, 190, band_h)[:, None]
    grad[..., 0], grad[..., 1], grad[..., 2] = 6, 30, 18
    grad[..., 3] = ramp
    overlay.paste(Image.fromarray(grad, "RGBA"), (0, image.height - band_h))
    image = Image.alpha_composite(image, overlay)
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_file(True), int(0.15 * DPI))
    draw.text(
        (int(0.16 * DPI), image.height - int(0.36 * DPI)),
        text,
        font=font,
        fill=(255, 255, 255, 255),
    )
    return image


def prepare_assets(metrics: dict) -> dict[str, Path]:
    """Generate every raster used by the deck and return their paths."""
    SLIDE_IMG.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}

    def save(name: str, image: Image.Image) -> None:
        path = SLIDE_IMG / name
        image.save(path, optimize=True)
        out[name] = path

    size = (1920, 1080)
    save("bg-a.png", gradient(size, "FFFFFF", "E7F1E2", 0.55))
    save("bg-b.png", gradient(size, "F7FBF4", "D9EAD2", 0.7))
    save("bg-capa.png", forest_watermark(gradient(size, "FFFFFF", "E1EFD9", 0.75), 0.2))
    logo = keyed_logo((110, 170, 1590, 705), drop_below=440, cutoff=226)
    save("logo-capa.png", logo)

    panel = gradient((int(6.0 * DPI), int(5.0 * DPI)), "0B3D24", "17693D", 0.85)
    save("panel-dark.png", rounded(panel, 0.18))
    term = gradient((int(6.0 * DPI), int(3.9 * DPI)), "08301D", "12492C", 0.9)
    save("panel-terminal.png", rounded(term, 0.16))

    photos = {
        "foto-acai.png": (
            "acai-paneiros.jpg",
            (0.5, 0.5),
            "Açaí em paneiros, Igarapé-Miri (PA)",
        ),
        "foto-castanha.png": (
            "castanha-extracao.jpg",
            (0.5, 0.35),
            "Extração de castanha-do-brasil",
        ),
        "foto-cupuacu.png": ("cupuacu.jpg", (0.5, 0.5), "Cupuaçu"),
        "foto-murumuru.png": ("murumuru.jpg", (0.5, 0.5), "Sementes de murumuru"),
    }
    for name, (source, focus, caption) in photos.items():
        image = caption_band(cover_crop(PRODUCTS / source, 2.95, 2.65, focus), caption)
        save(name, rounded(image, 0.14))
    save(
        "foto-castanha-larga.png",
        rounded(
            cover_crop(PRODUCTS / "castanha-extracao.jpg", 5.6, 3.15, (0.5, 0.4)), 0.14
        ),
    )

    def screen(source: str, name: str, box=None) -> None:
        image = Image.open(EVID / source).convert("RGB")
        if box:
            image = image.crop(box)
        save(name, rounded(image, 0.06))

    tile = (1280, 492)
    canvas_color = "#" + TINT

    def onto_canvas(source: str, name: str, width: int) -> None:
        image = Image.open(EVID / source).convert("RGB")
        scaled = image.resize(
            (width, int(image.height * width / image.width)), Image.LANCZOS
        )
        canvas = Image.new("RGB", tile, canvas_color)
        canvas.paste(
            scaled, ((tile[0] - scaled.width) // 2, (tile[1] - scaled.height) // 2)
        )
        save(name, rounded(canvas, 0.05))

    screen("01-painel.png", "demo-1-painel.png", (0, 0, tile[0], tile[1]))
    onto_canvas("03-registrar-confirmado.png", "demo-2-confirmado.png", 1040)
    screen(
        "04-consultar-lote.png", "demo-3-consulta.png", (0, 150, tile[0], 150 + tile[1])
    )
    onto_canvas("05-rejeicao-1.png", "demo-4-rejeicao.png", 1040)
    screen("10-adulteracao-detectada.png", "lab.png", (0, 0, 1280, 900))

    # Report pages for the bonus slide
    if REPORT_PDF.exists():
        text = subprocess.run(
            ["pdftotext", "-layout", str(REPORT_PDF), "-"],
            capture_output=True,
            text=True,
            check=False,
        ).stdout.split("\f")
        wanted = {
            "relatorio-capa.png": 1,
            "relatorio-arquitetura.png": next(
                (
                    i + 1
                    for i, t in enumerate(text)
                    if i > 2 and re.search(r"^\s*4 Arquitetura", t, re.M)
                ),
                5,
            ),
            "relatorio-cenarios.png": next(
                (
                    i + 1
                    for i, t in enumerate(text)
                    if i > 2 and "Apêndice A - Cenários de teste" in t
                ),
                15,
            ),
        }
        for name, page in wanted.items():
            target = SLIDE_IMG / name.replace(".png", "")
            subprocess.run(
                [
                    "pdftoppm",
                    "-png",
                    "-r",
                    "110",
                    "-f",
                    str(page),
                    "-l",
                    str(page),
                    "-singlefile",
                    str(REPORT_PDF),
                    str(target),
                ],
                check=False,
            )
            out[name] = SLIDE_IMG / name
    return out


# --------------------------------------------------------------- deck helpers
class Deck:
    """Helper around python-pptx with the deck's typography, header and footer."""

    def __init__(self, assets: dict[str, Path]) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.assets = assets
        self.number = 0

    def slide(self, background: str = "bg-a.png", chrome: bool = True):
        """Add a slide with a gradient background and, optionally, header and footer."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.shapes.add_picture(
            str(self.assets[background]),
            0,
            0,
            self.prs.slide_width,
            self.prs.slide_height,
        )
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)
        self.number += 1
        if chrome:
            self.picture(slide, LOGO, 0.7, 0.3, h=0.5)
            uea_h = 0.5
            uea_w = uea_h * 900 / 385
            self.picture(slide, UEA, 13.333 - 0.7 - uea_w, 0.3, h=uea_h)
            self.text(slide, 0.7, 7.02, 6.0, 0.3, COURSE, 10.5, color=MUTED)
            self.text(
                slide,
                11.93,
                7.02,
                0.7,
                0.3,
                str(self.number),
                10.5,
                color=MUTED,
                align=PP_ALIGN.RIGHT,
            )
        return slide

    @staticmethod
    def rect(
        slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None
    ):
        """Add a flat filled shape (no theme shadow)."""
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        style = shp._element.find(qn("p:style"))
        if style is not None:
            shp._element.remove(style)
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius
        return shp

    @staticmethod
    def text(
        slide,
        x,
        y,
        w,
        h,
        content,
        size,
        bold=False,
        color=INK,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,
        spacing=None,
        line_spacing=None,
        after=0,
        font=FONT,
    ):
        """Add a text box. ``content`` is a string or a list of paragraphs."""
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = (
            frame.margin_bottom
        ) = 0
        frame.vertical_anchor = anchor
        for index, item in enumerate(
            content if isinstance(content, list) else [content]
        ):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_after = Pt(after)
            if line_spacing:
                paragraph.line_spacing = line_spacing
            run = paragraph.add_run()
            run.text = item
            run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
            run.font.color.rgb = rgb(color)
            if spacing:
                run._r.get_or_add_rPr().set("spc", str(spacing))
        return box

    @staticmethod
    def picture(slide, path, x, y, w=None, h=None):
        """Place an image keeping its aspect ratio."""
        return slide.shapes.add_picture(
            str(path),
            Inches(x),
            Inches(y),
            Inches(w) if w else None,
            Inches(h) if h else None,
        )

    @staticmethod
    def line(slide, x1, y1, x2, y2, color=MUTED, width=1.25, dash=False, arrow=False):
        """Add a straight connector segment, optionally dashed or with an arrow head."""
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        connector.line.color.rgb = rgb(color)
        connector.line.width = Pt(width)
        if dash:
            connector.line.dash_style = 4
        if arrow:
            ln = connector.line._get_or_add_ln()
            tail = ln.makeelement(
                qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}
            )
            ln.append(tail)
        return connector

    @staticmethod
    def notes(slide, text: str) -> None:
        """Set the speaker notes."""
        slide.notes_slide.notes_text_frame.text = text


def overline(deck: Deck, slide, x, y, label, color=GREEN, w=4.0):
    """Small spaced caps label above a block."""
    deck.text(slide, x, y, w, 0.3, label.upper(), 11.5, True, color, spacing=160)


def title(deck: Deck, slide, text, x=0.7, y=1.05, w=11.9, size=32):
    """Slide title, left aligned."""
    deck.text(slide, x, y, w, 1.2, text, size, True, INK, line_spacing=0.95)


def pill(deck: Deck, slide, x, y, label, w=0.95):
    """Small filled tag used to mark bonus slides."""
    deck.rect(slide, x, y, w, 0.3, GREEN, MSO_SHAPE.ROUNDED_RECTANGLE, None, 0.5)
    deck.text(
        slide,
        x,
        y,
        w,
        0.3,
        label,
        10.5,
        True,
        "FFFFFF",
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
        spacing=120,
    )


# ------------------------------------------------------------------- slides
def build(metrics: dict, assets: dict[str, Path]) -> Deck:
    """Assemble the deck."""
    deck = Deck(assets)
    scen = metrics["scenarios"]
    total = scen["V"] + scen["I"] + scen["A"]
    pages = 21
    if REPORT_PDF.exists():
        info = subprocess.run(
            ["pdfinfo", str(REPORT_PDF)], capture_output=True, text=True, check=False
        ).stdout
        found = re.search(r"Pages:\s+(\d+)", info)
        pages = int(found.group(1)) if found else pages

    # 1. Cover
    s = deck.slide("bg-capa.png", chrome=False)
    deck.picture(s, assets["logo-capa.png"], (13.333 - 7.6) / 2, 1.05, w=7.6)
    deck.text(
        s,
        1.5,
        4.5,
        10.333,
        0.6,
        "Da origem ao consumidor, com transparência e integridade.",
        24,
        False,
        GREEN,
        PP_ALIGN.CENTER,
    )
    deck.text(
        s,
        1.5,
        5.3,
        10.333,
        0.4,
        "Ana Beatriz Maciel Nunes  |  Fernando Luiz da Silva Freire  |  Juliana Ballin Lima",
        14,
        False,
        MUTED,
        PP_ALIGN.CENTER,
    )
    deck.notes(
        s,
        "Abertura, 10 segundos. Apresentar o projeto e a equipe. Sugestão de fala: Ana.",
    )

    # 2. Context
    s = deck.slide()
    title(deck, s, "A bioeconomia amazônica vive de cadeias longas.", w=5.6, size=30)
    deck.text(
        s,
        0.7,
        3.05,
        5.4,
        1.9,
        "Açaí, castanha, cupuaçu e murumuru saem de comunidades extrativistas e passam por beneficiamento, transporte e distribuição até chegar ao consumidor.",
        16,
        False,
        MUTED,
        line_spacing=1.15,
    )
    deck.text(
        s,
        0.7,
        5.0,
        5.4,
        1.0,
        "Cada etapa fica com uma organização diferente.",
        20,
        True,
        GREEN,
        line_spacing=1.1,
    )
    positions = [
        ("foto-acai.png", 6.55, 1.2),
        ("foto-castanha.png", 9.65, 1.2),
        ("foto-cupuacu.png", 6.55, 4.0),
        ("foto-murumuru.png", 9.65, 4.0),
    ]
    for name, x, y in positions:
        deck.picture(s, assets[name], x, y, w=2.95)
    deck.text(
        s,
        6.55,
        6.72,
        6.1,
        0.25,
        "Fotos: Wikimedia Commons, licenças livres. Créditos em docs/assets/produtos.",
        9,
        False,
        MUTED,
    )
    deck.notes(
        s,
        "Bloco 1: 2 minutos para contexto, problema, dados e solução. Contexto: produtos da bioeconomia passam por várias organizações antes do consumidor. Fotos do Wikimedia Commons, créditos em docs/assets/produtos/CREDITOS.md. Sugestão de fala: Ana.",
    )

    # 3. Problem
    s = deck.slide()
    overline(deck, s, 0.7, 1.15, "Problema")
    deck.text(
        s,
        0.7,
        1.5,
        6.0,
        2.0,
        "Cada etapa da cadeia guarda o seu próprio registro.",
        32,
        True,
        INK,
        line_spacing=0.95,
    )
    deck.text(
        s,
        0.7,
        3.4,
        5.6,
        1.4,
        "O consumidor não tem como saber qual versão é a verdadeira, e quem administra o sistema pode reescrever o passado.",
        16,
        False,
        MUTED,
        line_spacing=1.15,
    )
    overline(deck, s, 0.7, 5.0, "Objetivo")
    deck.text(
        s,
        0.7,
        5.35,
        5.6,
        1.4,
        "Registrar cada etapa do lote de forma verificável, com regras que impedem atalhos.",
        19,
        True,
        GREEN,
        line_spacing=1.1,
    )
    top, box_h, gap = 1.35, 0.9, 0.25
    for index, name in enumerate(
        ["Produtor", "Beneficiador", "Transportador", "Distribuidor"]
    ):
        y = top + index * (box_h + gap)
        deck.rect(
            s, 7.3, y, 2.7, box_h, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.12
        )
        deck.text(s, 7.55, y + 0.15, 2.3, 0.35, name, 16, True, INK)
        deck.text(s, 7.55, y + 0.52, 2.3, 0.3, "registro próprio", 12, False, MUTED)
        deck.line(s, 10.0, y + box_h / 2, 10.45, y + box_h / 2, MUTED, 1.25, True)
    first, last = top + box_h / 2, top + 3 * (box_h + gap) + box_h / 2
    deck.line(s, 10.45, first, 10.45, last, MUTED, 1.25, True)
    mid = (first + last) / 2
    deck.line(s, 10.45, mid, 10.85, mid, MUTED, 1.25, True)
    deck.rect(
        s, 10.85, mid - 0.7, 1.75, 1.4, FOREST, MSO_SHAPE.ROUNDED_RECTANGLE, None, 0.1
    )
    deck.text(
        s,
        10.85,
        mid - 0.55,
        1.75,
        0.75,
        "?",
        40,
        True,
        LIME,
        PP_ALIGN.CENTER,
        MSO_ANCHOR.MIDDLE,
    )
    deck.text(
        s,
        10.85,
        mid + 0.2,
        1.75,
        0.4,
        "Consumidor",
        13,
        True,
        "FFFFFF",
        PP_ALIGN.CENTER,
    )
    deck.notes(
        s,
        "Problema: os registros ficam separados em cada organização e ninguém prova qual é o verdadeiro. Objetivo: registrar cada etapa de forma verificável, com regras. Sugestão de fala: Ana.",
    )

    # 4. Data
    s = deck.slide("bg-b.png")
    title(
        deck,
        s,
        "Só o açaí extrativo movimentou R$ 853 milhões em um ano.",
        w=11.9,
        size=30,
    )
    stats = [
        ("238,9 mil t", "de açaí extrativo colhidas no Brasil em 2023"),
        ("R$ 853,1 milhões", "valor da produção de açaí extrativo"),
        ("70,2%", "vêm do Pará, com 167,6 mil toneladas"),
    ]
    for index, (number, label) in enumerate(stats):
        y = 2.05 + index * 1.25
        deck.text(s, 0.7, y, 6.0, 0.75, number, 40, True, GREEN)
        deck.text(s, 0.7, y + 0.72, 5.8, 0.4, label, 14, False, MUTED)
    deck.picture(s, assets["foto-castanha-larga.png"], 7.0, 2.15, w=5.6)
    deck.text(
        s,
        7.0,
        5.45,
        5.6,
        0.9,
        "Castanha-do-pará: 35,4 mil toneladas e R$ 172,3 milhões. O Amazonas lidera, com 11,3 mil toneladas.",
        14,
        False,
        INK,
        line_spacing=1.1,
    )
    deck.text(
        s,
        0.7,
        5.85,
        6.0,
        0.7,
        "Quanto maior o valor, mais importa provar de onde o produto veio.",
        15,
        True,
        INK,
        line_spacing=1.05,
    )
    deck.text(
        s,
        0.7,
        6.68,
        11.9,
        0.25,
        "Fonte: IBGE, Produção da Extração Vegetal e da Silvicultura (PEVS) 2023. Foto: Wikimedia Commons, CC BY-SA 4.0.",
        9,
        False,
        MUTED,
    )
    deck.notes(
        s,
        "Dados oficiais do IBGE, PEVS 2023 (publicada em 2024): açaí extrativo com 238,9 mil toneladas e R$ 853,1 milhões, Pará com 167,6 mil toneladas (70,2%); castanha-do-pará com 35,4 mil toneladas e R$ 172,3 milhões, Amazonas liderando com 11,3 mil toneladas. Atenção: são números do extrativismo, não da produção cultivada. A frase final é a nossa conclusão, não um dado do IBGE. Sugestão de fala: Ana.",
    )

    # 5. Solution
    s = deck.slide()
    title(
        deck,
        s,
        "A solução: um histórico que ninguém reescreve sozinho.",
        w=11.9,
        size=30,
    )
    blocks = [
        ("Bloco 20", "0000…b3f1", "0000…7e5a"),
        ("Bloco 21", "0000…c4d9", "0000…b3f1"),
        ("Bloco 22", "0000…3a80", "0000…c4d9"),
    ]
    for index, (name, own, prev) in enumerate(blocks):
        x = 0.7 + index * 4.2
        deck.rect(
            s, x, 2.15, 3.5, 1.75, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.08
        )
        deck.text(s, x + 0.3, 2.35, 2.9, 0.35, name, 17, True, GREEN)
        deck.text(
            s, x + 0.3, 2.9, 2.9, 0.3, "hash        " + own, 12.5, False, INK, font=MONO
        )
        deck.text(
            s,
            x + 0.3,
            3.3,
            2.9,
            0.3,
            "anterior    " + prev,
            12.5,
            False,
            MUTED,
            font=MONO,
        )
        if index < 2:
            deck.line(s, x + 3.5, 3.02, x + 4.2, 3.02, GREEN, 2.0, arrow=True)
    deck.text(
        s,
        0.7,
        4.05,
        6.0,
        0.3,
        "Ilustração: o hash anterior repete o hash do bloco que o precede.",
        10.5,
        False,
        MUTED,
    )
    for index, (label, body) in enumerate(
        [
            (
                "Assinada",
                "Cada operação leva a assinatura digital da carteira responsável.",
            ),
            (
                "Encadeada",
                "Cada bloco guarda o hash do anterior. Alterar um bloco quebra a cadeia dali em diante.",
            ),
            (
                "Validada",
                "O contrato inteligente aplica perfis e a ordem das etapas antes de gravar.",
            ),
        ]
    ):
        x = 0.7 + index * 4.2
        overline(deck, s, x, 4.75, label, GREEN, 3.5)
        deck.text(s, x, 5.1, 3.6, 1.3, body, 15, False, INK, line_spacing=1.1)
    deck.text(
        s,
        0.7,
        6.5,
        11.9,
        0.35,
        "Faz sentido porque há vários participantes que não confiam plenamente uns nos outros.",
        13,
        False,
        MUTED,
    )
    deck.notes(
        s,
        "Solução: cada etapa vira uma transação assinada, validada pelo contrato e gravada em um bloco encadeado por hash. Contraste com um banco de dados: quem administra pode alterar e apagar, e a auditoria depende do próprio dono do dado. Ressalva: a blockchain garante que o registro não muda depois de gravado, não que ele seja verdadeiro na origem. Sugestão de fala: Ana.",
    )

    # 6. Architecture
    s = deck.slide()
    title(
        deck, s, "A carteira assina, o nó valida e o contrato decide.", w=11.5, size=30
    )
    deck.rect(s, 0.7, 1.9, 7.5, 4.45, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.03)
    deck.picture(s, DIAG / "arquitetura-slide.png", 0.85, 2.0, w=7.2)
    for offset, (label, body) in enumerate(
        [
            (
                "Navegador e gateway",
                "A interface nunca fala direto com o nó. O gateway assina com a carteira ativa.",
            ),
            (
                "Nó da blockchain",
                "Confere assinatura, nonce e data, minera com prova de trabalho e grava o bloco.",
            ),
            (
                "Contrato inteligente",
                "Perfis, ordem das etapas e validação das entradas, em Python.",
            ),
        ]
    ):
        y = 2.05 + offset * 1.55
        overline(deck, s, 8.6, y, label, GREEN, 3.9)
        deck.text(s, 8.6, y + 0.35, 4.0, 1.0, body, 15, False, INK, line_spacing=1.1)
    deck.notes(
        s,
        "Bloco 2: 2 minutos, junto com o próximo slide. Mostrar o caminho de uma operação: interface, gateway com carteira, nó, contrato e cadeia. Relacionar com o modelo de DApp da Aula 01. Sugestão de fala: Fernando.",
    )

    # 7. Contract
    s = deck.slide()
    title(deck, s, "O contrato só deixa o lote andar em ordem.", w=11.5, size=30)
    deck.rect(
        s, 0.7, 1.85, 11.9, 3.0, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.03
    )
    deck.picture(s, DIAG / "maquina-de-estados.png", 0.85, 1.9, w=11.6)
    for index, (label, body) in enumerate(
        [
            (
                "Quem pode",
                "Cada etapa exige um perfil: produtor, beneficiador, transportador ou distribuidor.",
            ),
            (
                "Em que ordem",
                "Pular ou repetir uma etapa é rejeitado. O lote finalizado não muda mais.",
            ),
            (
                "O que fica na cadeia",
                "Origem, etapas, responsáveis e o hash dos documentos. Os arquivos ficam de fora.",
            ),
        ]
    ):
        x = 0.7 + index * 4.2
        overline(deck, s, x, 5.05, label, GREEN, 3.6)
        deck.text(s, x, 5.4, 3.7, 1.3, body, 15, False, INK, line_spacing=1.1)
    deck.notes(
        s,
        "Explicar perfis, máquina de estados e regras de custódia: só o distribuidor designado confirma o recebimento. Tudo é validado antes de alterar o estado, como um revert do Solidity. Sugestão de fala: Fernando.",
    )

    # 8. Demonstration
    s = deck.slide("bg-b.png")
    title(deck, s, "Demonstração na blockchain local", size=28, w=8)
    deck.text(s, 9.4, 1.15, 3.2, 0.4, "4 minutos", 14, False, MUTED, PP_ALIGN.RIGHT)
    captions = [
        "1  Blockchain em execução",
        "2  Operação confirmada no bloco",
        "3  Registro e estado do lote",
        "4  Operação sem permissão rejeitada",
    ]
    tile_w = 5.5
    tile_h = tile_w * 492 / 1280
    for index, (name, caption) in enumerate(
        zip(
            [
                "demo-1-painel.png",
                "demo-2-confirmado.png",
                "demo-3-consulta.png",
                "demo-4-rejeicao.png",
            ],
            captions,
        )
    ):
        col, row = index % 2, index // 2
        x, y = 0.7 + col * (tile_w + 0.4), 1.75 + row * (tile_h + 0.6)
        deck.picture(s, assets[name], x, y, w=tile_w)
        deck.text(
            s,
            x,
            y + tile_h + 0.08,
            tile_w,
            0.3,
            caption,
            13,
            True,
            GREEN if index < 3 else RED,
        )
    deck.notes(
        s,
        "Bloco 3: 4 minutos, feito ao vivo na máquina da equipe. Os quatro passos exigidos: 1) mostrar a blockchain local em execução (make docker-up ou make start); 2) realizar uma operação pela interface e mostrar a confirmação com hash e bloco; 3) consultar o registro e mostrar a mudança de estado do lote; 4) rejeição de operação inválida ou sem permissão, usando os cenários da aba Registrar. Se a máquina falhar, usar o vídeo docs/slides/video/demonstracao.webm. Sugestão: Juliana executa os passos 1 e 2, Fernando o passo 3 e Ana o passo 4.",
    )

    # 9. Tests
    s = deck.slide()
    title(deck, s, "Testado no que funciona e no que deve falhar.", w=11.5, size=30)
    for index, (number, label) in enumerate(
        [
            (str(metrics["passed"]), "testes automatizados aprovados"),
            (f"{metrics['coverage']}%", "de cobertura do código"),
            (f"{scen['ok']} de {total}", "cenários com o resultado esperado"),
        ]
    ):
        y = 2.05 + index * 1.55
        deck.text(s, 0.7, y, 4.0, 0.9, number, 46, True, GREEN)
        deck.text(s, 0.7, y + 0.85, 4.0, 0.4, label, 14, False, MUTED)
    data = CategoryChartData()
    data.categories = [
        "Operações válidas",
        "Entradas inválidas ou sem permissão",
        "Integridade da cadeia",
    ]
    data.add_series("Cenários", (scen["V"], scen["I"], scen["A"]))
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(5.3),
        Inches(2.05),
        Inches(7.3),
        Inches(3.1),
        data,
    ).chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name, chart.font.size = FONT, Pt(13)
    chart.font.color.rgb = rgb(INK)
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    plot.data_labels.font.size, plot.data_labels.font.bold = Pt(14), True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = rgb(GREEN)
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.reverse_order = True
    chart.category_axis.format.line.color.rgb = rgb(LINE)
    chart.category_axis.tick_labels.font.size = Pt(13)
    deck.text(
        s,
        5.3,
        5.45,
        7.3,
        1.2,
        [
            "Toda entrada inválida ou sem permissão foi rejeitada sem criar bloco.",
            "A adulteração de um bloco foi detectada, com e sem refazer a prova de trabalho.",
        ],
        14,
        False,
        MUTED,
        after=6,
        line_spacing=1.1,
    )
    deck.notes(
        s,
        f"Bloco 4, parte 1. {metrics['passed']} testes pytest, cobertura de {metrics['coverage']}%, Pylint {metrics['pylint']} de 10, Bandit e pip-audit sem achados. Black Duck não foi executado por indisponibilidade da ferramenta. {total} cenários reais com resultado esperado versus obtido. Sugestão de fala: Juliana.",
    )

    # 10. Limits and conclusion
    s = deck.slide("bg-b.png")
    overline(deck, s, 0.7, 1.3, "Limitações")
    deck.text(
        s,
        0.7,
        1.75,
        5.2,
        4.6,
        [
            "Um único nó: a validação detecta a fraude, mas não há rede para recusá-la.",
            "Carteiras de demonstração guardadas no servidor, sem MetaMask.",
            "A blockchain garante o registro, não a verdade da informação na origem.",
        ],
        17,
        False,
        INK,
        after=18,
        line_spacing=1.1,
    )
    deck.picture(s, assets["panel-dark.png"], 6.6, 1.3, w=6.0)
    overline(deck, s, 7.1, 1.75, "Conclusão", LIME)
    deck.text(
        s,
        7.1,
        2.2,
        5.0,
        2.4,
        "Uma blockchain local em Python já sustenta a rastreabilidade do lote, do produtor ao consumidor.",
        26,
        True,
        "FFFFFF",
        line_spacing=1.02,
    )
    deck.text(
        s,
        7.1,
        4.75,
        5.0,
        1.1,
        "Próximos passos: vários nós com consenso, carteira no navegador e QR Code por lote.",
        15,
        False,
        PALE,
        line_spacing=1.1,
    )
    deck.text(
        s, 0.7, 6.55, 6.0, 0.3, "github.com/JulianaBallin/EcoOrigem", 13, False, GREEN
    )
    deck.notes(
        s,
        "Bloco 4, parte 2: limitações e conclusão, dentro dos 2 minutos. Ser direto sobre o que o protótipo não faz. Encerrar com os próximos passos. Depois deste slide, os quatro slides de bônus servem para perguntas. Sugestão de fala: Juliana.",
    )

    # 11. Bonus overview
    s = deck.slide("bg-b.png")
    pill(deck, s, 0.7, 1.1, "BÔNUS")
    title(deck, s, "O que entregamos além do pedido", y=1.55, size=30)
    extras = [
        (
            "01",
            "Relatório técnico",
            f"{pages} páginas em DOCX e PDF, com diagramas e apêndices",
        ),
        (
            "02",
            "Laboratório de integridade",
            "Simula a adulteração e mostra a validação detectando",
        ),
        (
            "03",
            f"{total} cenários de teste",
            "Resultado esperado e obtido, executados de verdade",
        ),
        (
            "04",
            "Docker e Makefile",
            "Uma linha sobe nó, contrato e interface, com menu interativo",
        ),
        (
            "05",
            "Vídeo de apoio",
            "Gravação da demonstração para o caso de falha técnica",
        ),
        (
            "06",
            "Documentação",
            "README, referência do contrato e roteiro da apresentação",
        ),
    ]
    for index, (num, head, body) in enumerate(extras):
        col, row = index % 2, index // 2
        x, y = 0.7 + col * 6.2, 2.55 + row * 1.35
        deck.text(s, x, y, 0.9, 0.7, num, 30, True, "9BC48A")
        deck.text(s, x + 0.95, y + 0.02, 4.7, 0.4, head, 18, True, INK)
        deck.text(
            s, x + 0.95, y + 0.48, 4.7, 0.7, body, 13.5, False, MUTED, line_spacing=1.1
        )
    deck.notes(
        s,
        "Slides de bônus, usar se sobrar tempo ou nas perguntas. Resumo de tudo que foi feito além do exigido pelo professor.",
    )

    # 12. Bonus report
    s = deck.slide()
    pill(deck, s, 0.7, 1.1, "BÔNUS")
    title(deck, s, "Um relatório técnico completo", y=1.55, w=5.4, size=30)
    deck.text(
        s,
        0.7,
        2.75,
        5.0,
        2.6,
        [
            f"{pages} páginas com problema, justificativa, arquitetura, implementação, testes, limitações e conclusão.",
            "Diagramas, tabelas de perfis e métodos, e a tabela completa dos cenários de teste no apêndice.",
        ],
        15,
        False,
        MUTED,
        after=12,
        line_spacing=1.15,
    )
    deck.text(s, 0.7, 5.5, 5.0, 0.5, "docs/relatorio", 13, False, GREEN, font=MONO)
    for index, (name, caption) in enumerate(
        [
            ("relatorio-capa.png", "Capa"),
            ("relatorio-arquitetura.png", "Arquitetura"),
            ("relatorio-cenarios.png", "Cenários de teste"),
        ]
    ):
        x = 6.1 + index * 2.3
        if name in assets and assets[name].exists():
            deck.rect(s, x - 0.03, 1.6, 2.1 + 0.06, 2.97 + 0.06, LINE)
            deck.picture(s, assets[name], x, 1.63, w=2.1)
        deck.text(s, x, 4.75, 2.1, 0.3, caption, 12, False, MUTED)
    deck.notes(
        s,
        "O relatório segue o padrão da UEA: capa em página inteira com logos, texto justificado, cabeçalho, rodapé e legendas nas figuras. É gerado por script (make report), o que mantém os números de testes sempre atualizados.",
    )

    # 13. Bonus integrity lab
    s = deck.slide("bg-b.png")
    pill(deck, s, 0.7, 1.1, "BÔNUS")
    title(
        deck,
        s,
        "Alterar um dado gravado não passa despercebido",
        y=1.55,
        w=5.3,
        size=28,
    )
    deck.text(
        s,
        0.7,
        3.3,
        5.0,
        3.0,
        [
            "Um campo de um bloco é alterado em memória.",
            "A validação aponta o bloco e o motivo, como raiz de Merkle e assinatura divergentes.",
            "O nó bloqueia novas operações até a cadeia ser restaurada.",
        ],
        15,
        False,
        MUTED,
        after=12,
        line_spacing=1.15,
    )
    deck.picture(s, assets["lab.png"], 6.3, 1.5, w=6.3)
    deck.notes(
        s,
        "Laboratório de integridade da interface: adulterar sem refazer hashes, ou refazendo a prova de trabalho do bloco (detectado pela quebra de encadeamento com o bloco seguinte). O arquivo em disco não é alterado, e o botão Restaurar recarrega a cadeia.",
    )

    # 14. Bonus reproducible run
    s = deck.slide()
    pill(deck, s, 0.7, 1.1, "BÔNUS")
    title(deck, s, "Execução reproduzível", y=1.55, w=5.3, size=30)
    deck.text(
        s,
        0.7,
        2.6,
        5.0,
        3.6,
        [
            "Docker Compose sobe o nó, implanta o contrato e inicia a interface, nessa ordem.",
            "O Makefile cria o ambiente virtual e reúne testes, lint, cenários e documentação.",
            "O vídeo de apoio fica em docs/slides/video.",
        ],
        15,
        False,
        MUTED,
        after=12,
        line_spacing=1.15,
    )
    deck.picture(s, assets["panel-terminal.png"], 6.6, 1.6, w=6.0)
    lines = [
        "$ make docker-up",
        "$ make docker-seed",
        "$ make menu",
        "$ make test",
        "$ make scenarios",
        "$ make report slides",
    ]
    for index, line in enumerate(lines):
        deck.text(
            s, 7.05, 2.0 + index * 0.55, 5.2, 0.4, line, 16, False, "D7F0C8", font=MONO
        )
    deck.notes(
        s,
        "Reprodutibilidade: um clone limpo cria o venv pelo Makefile e passa nos 248 testes. O serviço deploy do Compose garante que o contrato esteja implantado antes de a interface subir.",
    )
    return deck


def main() -> int:
    """Build the PPTX and convert it to PDF."""
    if not METRICS.exists():
        print("Execute make report antes: faltam as métricas medidas.")
        return 1
    metrics = json.loads(METRICS.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    deck = build(metrics, prepare_assets(metrics))
    deck.prs.save(str(OUT))
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUT_DIR),
            str(OUT),
        ],
        check=False,
        capture_output=True,
    )
    print(f"Apresentação gerada: {OUT} e {OUT.with_suffix('.pdf')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
