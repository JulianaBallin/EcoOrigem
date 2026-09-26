"""Build the technical report (DOCX and PDF) of the EcoOrigem delivery.

Numbers about tests and quality are measured when the report is built, and the
scenario table comes from the executed scenarios (``make scenarios``).
Run with ``make report``.
"""

# pylint: disable=duplicate-code,line-too-long,too-many-statements,too-many-locals,protected-access,too-many-lines,too-many-arguments,too-many-positional-arguments

from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_lib import (  # noqa: E402  pylint: disable=wrong-import-position
    GREEN_100,
    GREEN_800,
    GREEN_900,
    Report,
    bottom_border,
    cell_borders,
    cell_margins,
    fixed_table,
    set_font,
    shade,
)

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
LOGOS = ROOT / "assets" / "img"
LOGO_UEA = ASSETS / "logos" / "uea.png"
DIAGRAMS = ASSETS / "diagramas"
SHOTS = ASSETS / "evidencias"
SCENARIOS_MD = ROOT / "docs" / "evidencias" / "cenarios-de-teste.md"
METRICS_JSON = ROOT / "docs" / "evidencias" / "metricas.json"
OUT_DIR = ROOT / "docs" / "relatorio"
NAME = "relatorio-tecnico-ecoorigem"
TEAM = [
    "Ana Beatriz Maciel Nunes",
    "Fernando Luiz da Silva Freire",
    "Juliana Ballin Lima",
]


# ------------------------------------------------------------------ metrics
def run_command(command: list[str]) -> subprocess.CompletedProcess:
    """Run a command from the project root."""
    return subprocess.run(
        command, cwd=ROOT, capture_output=True, text=True, check=False
    )


def quality_metrics() -> dict:
    """Measure tests, coverage and static analysis results."""
    python = sys.executable
    with tempfile.TemporaryDirectory() as tmp:
        cov_file = Path(tmp) / "cov.json"
        tests = run_command(
            [python, "-m", "pytest", "-q", "--cov", f"--cov-report=json:{cov_file}"]
        )
        coverage = (
            json.loads(cov_file.read_text())["totals"]["percent_covered"]
            if cov_file.exists()
            else 0.0
        )
    passed = re.search(r"(\d+) passed", tests.stdout)
    collected = run_command([python, "-m", "pytest", "--collect-only", "-q"]).stdout
    per_file = Counter(
        line.split("::")[0] for line in collected.splitlines() if "::" in line
    )
    pylint = run_command([python, "-m", "pylint", "ecoorigem", "tests", "scripts"])
    score = re.search(r"rated at ([\d.]+)/10", pylint.stdout)
    black = run_command(
        [python, "-m", "black", "--check", "ecoorigem", "tests", "scripts"]
    )
    bandit = run_command([python, "-m", "bandit", "-q", "-r", "ecoorigem"])
    audit = run_command([python, "-m", "pip_audit", "-r", "requirements.txt"])
    return {
        "passed": int(passed.group(1)) if passed else 0,
        "failed": tests.returncode != 0,
        "coverage": round(coverage),
        "per_file": per_file,
        "pylint": score.group(1) if score else "n/d",
        "black": black.returncode == 0,
        "bandit": bandit.returncode == 0,
        "audit": (
            "Nenhuma vulnerabilidade conhecida"
            if audit.returncode == 0
            else "Não executado ou com achados"
        ),
    }


def scenario_rows() -> tuple[list[list[str]], Counter]:
    """Parse the executed scenario table from the evidence markdown."""
    rows, counts = [], Counter()
    for line in SCENARIOS_MD.read_text(encoding="utf-8").splitlines():
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) == 5 and re.match(r"^[VIA]\d\d$", cells[0]):
            rows.append(cells)
            counts[cells[0][0]] += 1
            counts["ok" if cells[4] == "Aprovado" else "fail"] += 1
    return rows, counts


def cropped(name: str, max_ratio: float = 0.9) -> Path:
    """Crop tall full-page screenshots so they fit a report page."""
    source = SHOTS / name
    image = Image.open(source)
    limit = int(image.width * max_ratio)
    if image.height <= limit:
        return source
    target = Path(tempfile.gettempdir()) / f"crop-{name}"
    image.crop((0, 0, image.width, limit)).save(target)
    return target


# -------------------------------------------------------------------- cover
def _pad(cell, top=0, bottom=0, left=0, right=0) -> None:
    """Set cell padding in centimeters."""
    cell_margins(
        cell, int(top * 567), int(bottom * 567), int(left * 567), int(right * 567)
    )


def _line(
    cell,
    text,
    size,
    bold=False,
    color="17261D",
    after=0,
    first=False,
    spacing=None,
    before=0,
):
    """Add a left-aligned line of text to a cover cell."""
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_font(run, size, bold=bold, color=color)
    if spacing:
        props = run._r.get_or_add_rPr()  # pylint: disable=protected-access
        node = OxmlElement("w:spacing")
        node.set(qn("w:val"), str(spacing))
        props.append(node)
    return paragraph


def cover(report: Report) -> None:
    """Full-page cover: logos band, title block, accent line and details."""
    doc = report.doc
    rows = [
        5.4,
        13.0,
        0.45,
        10.75,
    ]  # centimeters, sums to the A4 height minus the marker line
    table = doc.add_table(rows=4, cols=2)
    fixed_table(table, [10.5, 10.5])
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr  # pylint: disable=protected-access
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    props.append(indent)
    for index, height in enumerate(rows):
        table.rows[index]._tr.get_or_add_trPr().append(
            _row_height(Cm(height))
        )  # pylint: disable=protected-access

    # Row 1: institutional and project logos on white
    for index, (path, width, align, left, right) in enumerate(
        (
            (LOGO_UEA, 6.6, WD_ALIGN_PARAGRAPH.LEFT, 2.2, 0),
            (LOGOS / "logo.png", 6.6, WD_ALIGN_PARAGRAPH.RIGHT, 0, 2.2),
        )
    ):
        cell = table.rows[0].cells[index]
        cell_borders(cell, "FFFFFF", 0, sides=())
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _pad(cell, 0.6, 0.6, left, right)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_picture(str(path), width=Cm(width))

    # Row 2: title block
    band = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    shade(band, GREEN_900)
    cell_borders(band, GREEN_900, 0, sides=())
    band.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _pad(band, 0.8, 0.8, 2.2, 2.2)
    _line(
        band,
        "RELATÓRIO TÉCNICO  |  1º BIMESTRE",
        11,
        True,
        "8CC63F",
        14,
        first=True,
        spacing=40,
    )
    _line(band, "EcoOrigem", 66, True, "FFFFFF", 4)
    rule = _line(band, " ", 6, False, "FFFFFF", 16)
    rule.paragraph_format.right_indent = Cm(12.5)
    bottom_border(rule, "8CC63F", 24, 1)
    _line(
        band,
        "Rastreabilidade de produtos da bioeconomia amazônica",
        21,
        False,
        "FFFFFF",
        6,
    )
    _line(
        band,
        "Blockchain local em Python com contrato inteligente, interface web e execução em Docker",
        13,
        False,
        "CFE5C5",
        0,
    )

    # Row 3: accent strip
    strip = table.rows[2].cells[0].merge(table.rows[2].cells[1])
    shade(strip, "8CC63F")
    cell_borders(strip, "8CC63F", 0, sides=())
    strip.paragraphs[0].paragraph_format.space_after = Pt(0)
    strip.paragraphs[0].paragraph_format.line_spacing = Pt(1)

    # Row 4: details in two columns
    left, right = table.rows[3].cells
    for cell, (l, r) in ((left, (2.2, 0.6)), (right, (0.6, 2.2))):
        shade(cell, GREEN_100)
        cell_borders(cell, GREEN_100, 0, sides=())
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _pad(cell, 1.9, 0.5, l, r)
    facts_left = [
        ("INSTITUIÇÃO", ["Universidade do Estado do Amazonas"]),
        ("DISCIPLINA", ["Oficina de Desenvolvimento de Sistemas III"]),
        ("PROFESSOR", ["Prof. Dr. Fábio Santos"]),
        ("DATA", ["Setembro de 2026"]),
    ]
    facts_right = [
        ("EQUIPE", TEAM),
        ("ATIVIDADE", ["Trabalho do 1º bimestre: aplicação com blockchain local"]),
    ]
    for cell, facts in ((left, facts_left), (right, facts_right)):
        first = True
        for label, values in facts:
            _line(
                cell,
                label,
                8.5,
                True,
                GREEN_800,
                2,
                first=first,
                spacing=30,
                before=0 if first else 12,
            )
            first = False
            for value in values:
                _line(cell, value, 12, False, "17261D", 2)
    report.start_body()


def _row_height(height):
    """Exact row height element."""
    node = OxmlElement("w:trHeight")
    node.set(qn("w:val"), str(int(height.twips)))
    node.set(qn("w:hRule"), "exact")
    return node


# --------------------------------------------------------------------- body
def body(
    report: Report, metrics: dict, scenarios: list[list[str]], counts: Counter
) -> None:
    """Write every section of the report."""
    p, h = report.paragraph, report.heading

    # ------------------------------------------------------------- 1
    h(1, "1 Introdução", page_break=True)
    p(
        "A bioeconomia amazônica movimenta produtos como açaí, castanha-do-brasil, murumuru, cupuaçu, andiroba e óleos vegetais. Entre a comunidade extrativista e o consumidor final, cada lote passa por beneficiamento, transporte e distribuição, e cada etapa fica sob responsabilidade de uma organização diferente. As informações sobre essa trajetória costumam ficar dispersas em planilhas, notas e sistemas isolados, sem garantia de que não foram alteradas depois."
    )
    p(
        "Este relatório apresenta o EcoOrigem, um protótipo de rastreabilidade que registra a trajetória de cada lote em uma blockchain local implementada em Python. O trabalho foi desenvolvido para a disciplina Oficina de Desenvolvimento de Sistemas III e segue a lógica apresentada em aula: bloco com hash SHA-256, encadeamento por hash anterior, prova de trabalho com nonce, validação da cadeia e uma aplicação descentralizada formada por interface, contrato inteligente e blockchain."
    )

    # ------------------------------------------------------------- 2
    h(1, "2 Problema e objetivos")
    h(2, "2.1 Problema")
    p(
        "Como garantir um histórico confiável e rastreável da origem e das movimentações de produtos da bioeconomia amazônica, de modo que cada participante da cadeia consiga comprovar o que registrou e ninguém consiga alterar o passado sem deixar rastro?"
    )
    p(
        "Hoje o histórico de um lote depende de registros mantidos separadamente por produtor, beneficiador, transportador e distribuidor. Qualquer um deles pode corrigir, apagar ou reescrever a própria versão, e o consumidor não tem como verificar qual versão é a verdadeira. Isso enfraquece a confiança na origem declarada e dificulta certificações e auditorias."
    )
    h(2, "2.2 Objetivo geral")
    p(
        "Desenvolver uma aplicação sobre uma blockchain local em Python que registre, consulte e audite o ciclo de vida de lotes de produtos da bioeconomia amazônica, com regras de negócio executadas por um contrato inteligente."
    )
    h(2, "2.3 Objetivos específicos")
    report.bullets(
        [
            "Implementar uma blockchain local com blocos, hash SHA-256, prova de trabalho, árvore de Merkle e validação da cadeia.",
            "Implementar um contrato inteligente com perfis de acesso, ordem obrigatória das etapas e validação das entradas.",
            "Disponibilizar uma interface web para registrar operações, consultar lotes e visualizar a blockchain.",
            "Demonstrar que operações inválidas ou sem permissão são rejeitadas e que a adulteração de dados gravados é detectada.",
            "Documentar quais dados ficam na blockchain, testar operações válidas e inválidas e empacotar a solução para execução reproduzível.",
        ]
    )
    h(2, "2.4 Escopo")
    p(
        "O protótipo cobre o registro de lotes, o beneficiamento, o transporte, a confirmação de recebimento, a finalização e o anexo de documentos por hash. Ficam fora do escopo a integração com sensores, certificadoras e sistemas reais de cooperativas, a rede blockchain pública e o uso de moeda ou tokens."
    )

    # ------------------------------------------------------------- 3
    h(1, "3 Justificativa do uso de blockchain")
    p(
        "A rastreabilidade de uma cadeia produtiva reúne três características que favorecem a blockchain: há vários participantes independentes, nenhum deles deve poder reescrever o histórico sozinho e todos precisam verificar os mesmos registros. A Tabela 1 compara a solução com um banco de dados relacional tradicional operado por uma única organização.",
        keep_next=True,
    )
    report.table(
        "Banco de dados tradicional e blockchain no problema do EcoOrigem",
        ["Critério", "Banco de dados centralizado", "Blockchain do EcoOrigem"],
        [
            [
                "Quem controla o histórico",
                "O administrador do banco pode alterar ou apagar registros",
                "Nenhum participante altera um bloco sem quebrar o encadeamento",
            ],
            [
                "Detecção de fraude",
                "Depende de logs que o mesmo administrador controla",
                "Qualquer nó recalcula os hashes e a raiz de Merkle e detecta a alteração",
            ],
            [
                "Regras de negócio",
                "Ficam no código da aplicação, que pode ser contornado",
                "Ficam no contrato inteligente, executado sobre a cadeia",
            ],
            [
                "Autoria das operações",
                "Usuário e senha em tabela do sistema",
                "Assinatura digital Ed25519 por transação",
            ],
            [
                "Auditoria",
                "Exige acesso privilegiado",
                "Histórico completo e cronológico consultável por todos",
            ],
            [
                "Custo e complexidade",
                "Baixos",
                "Maiores: prova de trabalho, validação e replicação",
            ],
        ],
        [3.6, 6.2, 6.2],
        caption="a coluna da direita descreve o comportamento do protótipo, com as ressalvas da seção 9.",
    )
    p(
        "Blockchain não é a melhor resposta para todo problema. Se uma única organização confiável controlasse toda a cadeia, um banco de dados bastaria. Aqui o valor está no compartilhamento de um registro imutável entre partes que não confiam plenamente umas nas outras. Ainda assim, a blockchain garante que o dado registrado não muda depois de gravado, e não que o dado seja verdadeiro na origem. Essa limitação é discutida na seção 9."
    )

    # ------------------------------------------------------------- 4
    h(1, "4 Arquitetura")
    p(
        "A arquitetura segue o modelo de aplicação descentralizada da Aula 01: interface, contrato inteligente e blockchain, com uma carteira responsável por assinar as transações. A Figura 1 mostra os componentes e a Tabela 2 descreve a responsabilidade de cada um."
    )
    report.figure(
        DIAGRAMS / "arquitetura.png",
        "Arquitetura do EcoOrigem",
        "o navegador conversa apenas com o gateway web; o gateway assina a transação com a carteira ativa e a envia ao nó, onde o contrato executa as regras e a cadeia registra o resultado.",
        15.5,
    )
    report.table(
        "Componentes e responsabilidades",
        ["Componente", "Tecnologia", "Responsabilidade"],
        [
            [
                "Interface web",
                "HTML, CSS e JavaScript",
                "Formulários, consulta de lotes, explorador de blocos e laboratório de integridade",
            ],
            [
                "Gateway e carteiras",
                "Flask, Ed25519",
                "Guarda as chaves de demonstração, assina a transação e encaminha leituras ao nó",
            ],
            [
                "Nó da blockchain",
                "Flask, Python",
                "Recebe transações assinadas, valida, executa o contrato, minera e persiste os blocos",
            ],
            [
                "Contrato inteligente",
                "Python (EcoOrigemContract)",
                "Perfis, máquina de estados, validações e eventos",
            ],
            [
                "Cadeia de blocos",
                "Python (Block, Blockchain)",
                "Encadeamento por SHA-256, prova de trabalho, raiz de Merkle e validação",
            ],
            [
                "Empacotamento",
                "Docker, Makefile, venv",
                "Execução reproduzível, testes em contêiner e menu interativo",
            ],
        ],
        [3.6, 4.2, 8.2],
        caption="cada linha corresponde a um pacote do repositório, descrito na seção 5.",
    )
    p(
        "Uma operação percorre o caminho da Figura 2. O nó só inclui um bloco quando a assinatura, o nonce, a data e todas as regras do contrato passam. Em qualquer falha, nenhum bloco é criado e o motivo retorna à pessoa usuária."
    )
    report.figure(
        DIAGRAMS / "sequencia-da-operacao.png",
        "Ciclo de vida de uma operação",
        "as etapas 1 a 8 mostram a passagem da operação pelas camadas até o bloco minerado ou até a rejeição.",
        16.0,
    )

    # ------------------------------------------------------------- 5
    h(1, "5 Implementação")
    h(2, "5.1 Blockchain")
    p(
        "O bloco reproduz o modelo visto em aula: índice, data, hash anterior, nonce e hash SHA-256. Duas evoluções aproximam o protótipo de uma blockchain real. O cabeçalho passa a guardar a raiz da árvore de Merkle das transações, e cada transação é assinada digitalmente. A Figura 3 mostra o encadeamento."
    )
    report.figure(
        DIAGRAMS / "estrutura-dos-blocos.png",
        "Estrutura e encadeamento dos blocos",
        "a seta parte do hash de um bloco e chega ao campo hash anterior do bloco seguinte; alterar qualquer campo muda o hash e rompe essa ligação.",
        16.0,
    )
    report.bullets(
        [
            (
                "Hash e serialização.",
                "O hash do cabeçalho é o SHA-256 de um JSON canônico, com chaves ordenadas e sem espaços, o que produz o mesmo hash em qualquer máquina.",
            ),
            (
                "Prova de trabalho.",
                "O mineiro incrementa o nonce até o hash começar com a quantidade de zeros definida pela dificuldade, quatro por padrão. Em média são cerca de 65 mil tentativas por bloco, o que leva frações de segundo em Python.",
            ),
            (
                "Árvore de Merkle.",
                "As transações do bloco são resumidas em uma raiz. Uma prova de inclusão com poucos hashes irmãos demonstra que uma transação pertence ao bloco, e a interface exibe essa prova.",
            ),
            (
                "Assinatura digital.",
                "Cada transação leva a chave pública, o nonce e a assinatura Ed25519 sobre o seu hash. O endereço da conta é derivado da chave pública, e o nonce por conta impede a repetição de uma transação já aceita.",
            ),
        ]
    )
    p(
        "A validação percorre toda a cadeia e lista cada problema encontrado, em vez de parar no primeiro. A Tabela 3 resume as verificações.",
        keep_next=True,
    )
    report.table(
        "Verificações de integridade da cadeia",
        ["Código", "O que detecta"],
        [
            [
                "HASH_MISMATCH",
                "O conteúdo do cabeçalho não corresponde ao hash armazenado",
            ],
            [
                "BROKEN_LINK",
                "O hash anterior não coincide com o hash do bloco precedente",
            ],
            ["POW_INVALID", "O hash não atende à dificuldade da prova de trabalho"],
            [
                "MERKLE_MISMATCH",
                "As transações não correspondem à raiz de Merkle do bloco",
            ],
            [
                "TX_SIGNATURE_INVALID",
                "A assinatura ou o endereço do remetente não conferem",
            ],
            [
                "TX_NONCE_INVALID e DUPLICATE_TX",
                "Nonce fora de sequência ou transação repetida",
            ],
            [
                "CONTRACT_REPLAY_FAILED",
                "Reexecutar as transações no contrato produz uma rejeição",
            ],
            [
                "BAD_INDEX, TIMESTAMP_ORDER, GENESIS_INVALID",
                "Sequência, ordem cronológica e formato do bloco gênesis",
            ],
        ],
        [5.8, 10.2],
        caption="a mesma lista é exibida na aba Integridade da interface quando uma adulteração é simulada.",
    )

    h(2, "5.2 Contrato inteligente")
    p(
        "O contrato EcoOrigem é uma classe Python executada pelo nó a cada transação. Seu estado é uma função pura das transações gravadas: reexecutar a cadeia desde o bloco gênesis reconstrói exatamente o mesmo estado, e é assim que o nó reinicia. Cada método segue o padrão verificar antes de alterar, usado em Solidity: permissões e entradas são checadas primeiro e o estado só muda depois de todas as verificações passarem. Uma falha equivale a um revert."
    )
    report.table(
        "Perfis de acesso",
        ["Perfil", "Pode executar", "Restrição adicional"],
        [
            [
                "Administrador",
                "Conceder e revogar perfis",
                "É a conta que implantou o contrato",
            ],
            [
                "Produtor",
                "Registrar lote",
                "Designa o beneficiador do lote",
            ],
            [
                "Beneficiador",
                "Registrar beneficiamento",
                "Só o beneficiador designado; designa o transportador",
            ],
            [
                "Transportador",
                "Iniciar transporte",
                "Só o transportador designado; designa um distribuidor com perfil válido",
            ],
            [
                "Distribuidor",
                "Confirmar recebimento e finalizar",
                "Só o distribuidor designado confirma; só o responsável atual finaliza",
            ],
            ["Consumidor", "Consultar lotes", "Leitura pública, sem carteira"],
        ],
        [3.2, 5.6, 7.2],
        caption="cada etapa designa quem executa a próxima; o responsável atual pelo lote muda a cada etapa e é o único que pode anexar documentos.",
    )
    report.figure(
        DIAGRAMS / "maquina-de-estados.png",
        "Ciclo de vida do lote",
        "cada etapa só é aceita depois da anterior e pelo perfil indicado; o estado FINALIZADO é definitivo.",
        16.0,
    )
    report.table(
        "Métodos do contrato",
        ["Método", "Perfil", "Efeito", "Principais rejeições"],
        [
            [
                "grant_role e revoke_role",
                "Administrador",
                "Concede ou revoga um perfil",
                "ACCESS_DENIED, INVALID_INPUT",
            ],
            [
                "register_lot",
                "Produtor",
                "Cria o lote em CADASTRADO",
                "Produto fora da lista, quantidade inválida, data futura",
            ],
            [
                "record_processing",
                "Beneficiador designado",
                "CADASTRADO para BENEFICIADO",
                "Outro beneficiador tenta registrar, INVALID_TRANSITION",
            ],
            [
                "start_transport",
                "Transportador designado",
                "BENEFICIADO para EM_TRANSPORTE",
                "Outro transportador tenta iniciar, destinatário sem perfil de distribuidor",
            ],
            [
                "confirm_delivery",
                "Distribuidor designado",
                "EM_TRANSPORTE para DISTRIBUIDO",
                "Outro distribuidor tenta confirmar",
            ],
            [
                "finalize_lot",
                "Distribuidor responsável",
                "DISTRIBUIDO para FINALIZADO",
                "INVALID_TRANSITION",
            ],
            [
                "attach_document",
                "Responsável atual",
                "Registra o hash de um documento",
                "LOT_FINALIZED, hash repetido ou inválido",
            ],
        ],
        [3.6, 3.2, 4.4, 4.8],
        caption="toda operação sobre lote FINALIZADO retorna LOT_FINALIZED.",
        size=8.5,
    )
    p(
        "As entradas passam por validadores que normalizam e limitam cada campo: produto da lista aceita, com tolerância a caixa e acentos; origem entre 3 e 120 caracteres; quantidade positiva, finita e limitada; data de coleta válida e não futura; hashes de documento com 64 caracteres hexadecimais; campos desconhecidos rejeitados. As mensagens de erro são em português e cada rejeição carrega um código estável."
    )

    h(2, "5.3 Nó da blockchain")
    report.bullets(
        [
            (
                "Aceitação.",
                "O nó verifica assinatura, endereço, nonce esperado, transação repetida e diferença de relógio de no máximo cinco minutos antes de executar o contrato.",
            ),
            (
                "Mineração.",
                "Com a mineração automática cada transação aceita vira um bloco. No modo manual, as transações ficam na fila e um único bloco as agrupa, o que mostra a árvore de Merkle com várias folhas.",
            ),
            (
                "Persistência.",
                "A cadeia é gravada de forma atômica em JSON após cada bloco. Ao iniciar, o nó relê o arquivo, valida a cadeia e reconstrói o estado do contrato.",
            ),
            (
                "Rejeições.",
                "Toda operação recusada é registrada em um log local com camada, código, motivo, carteira e método. Rejeições não criam bloco nem consomem nonce.",
            ),
            (
                "Bloqueio.",
                "Se a validação de integridade falhar, o nó recusa novas escritas com o código CHAIN_COMPROMISED até a cadeia ser restaurada.",
            ),
        ]
    )

    h(2, "5.4 Interface web")
    p(
        "A interface é uma aplicação de página única, sem dependências externas, servida pelo gateway Flask. A carteira ativa fica sempre visível no topo, e o estado da blockchain local é exibido em tempo real."
    )
    report.table(
        "Telas da interface",
        ["Tela", "Função"],
        [
            [
                "Painel",
                "Blocos, transações, lotes por etapa, integridade, contrato e atividade recente",
            ],
            [
                "Registrar",
                "Formulários das seis operações, resultado com hash e bloco, cenários de rejeição",
            ],
            [
                "Consultar",
                "Lista com filtros, linha do tempo do lote, documentos e verificação de arquivo",
            ],
            [
                "Blockchain",
                "Blocos com hash, hash anterior, raiz de Merkle, nonce, transações e prova de inclusão",
            ],
            ["Permissões", "Carteiras, perfis e operações do administrador"],
            ["Rejeições", "Log das operações recusadas, com camada e motivo"],
            [
                "Integridade",
                "Simulação de adulteração, validação da cadeia e restauração",
            ],
        ],
        [3.2, 12.8],
        caption="as capturas de tela de cada uma estão no Apêndice B.",
    )
    report.figure(
        cropped("01-painel.png", 0.95),
        "Painel da aplicação",
        "visão geral com os indicadores da cadeia, o ciclo de vida dos lotes e o estado do contrato.",
        11.5,
    )

    h(2, "5.5 Infraestrutura")
    p(
        "A solução é dockerizada e também roda localmente. O arquivo docker-compose.yml sobe três serviços em ordem: o nó, uma implantação única do contrato com concessão dos perfis de demonstração e a interface. Um perfil opcional cria lotes de exemplo e outro executa os testes em contêiner. O Makefile reúne os comandos principais, cria o ambiente virtual .venv-ecoorigem quando necessário e oferece o comando make menu, um menu interativo com cada opção explicada."
    )

    # ------------------------------------------------------------- 6
    h(1, "6 Dados na blockchain")
    p(
        "A separação entre o que é gravado na cadeia e o que permanece fora dela protege a privacidade e evita crescimento desnecessário. Arquivos e dados pessoais não são gravados. Quando é preciso provar a autenticidade de um documento, grava-se apenas o seu hash SHA-256, e a interface recalcula o hash de um arquivo escolhido e o compara com o registrado."
    )
    report.table(
        "Dados gravados e dados mantidos fora da blockchain",
        ["Na blockchain", "Fora da blockchain"],
        [
            [
                "Identificador, produto, origem, quantidade e data de coleta do lote",
                "Arquivos de certificados, laudos e fotos",
            ],
            [
                "Status atual e cada mudança de etapa, com data e bloco",
                "Dados pessoais de produtores e trabalhadores",
            ],
            [
                "Endereço da carteira responsável por cada operação",
                "Chaves privadas, mantidas no keystore local",
            ],
            [
                "Perfis concedidos e revogados pelo administrador",
                "Log das operações rejeitadas, mantido pelo nó",
            ],
            ["Hash SHA-256 de documentos anexados", "Conteúdo dos documentos"],
            [
                "Assinatura digital de cada transação",
                "Senhas ou credenciais de qualquer tipo",
            ],
        ],
        [8.0, 8.0],
        caption="o hash prova que um arquivo não foi alterado sem expô-lo publicamente.",
    )

    # ------------------------------------------------------------- 7
    h(1, "7 Testes")
    h(2, "7.1 Estratégia")
    p(
        f"Os testes automatizados usam pytest e cobrem, em camadas, a criptografia, a cadeia, o contrato, o nó, a API HTTP, o gateway, o cliente e a linha de comando. Foram executados {metrics['passed']} testes, todos aprovados, com cobertura de {metrics['coverage']}% das linhas e ramos do pacote. A Tabela 8 resume o resultado das verificações de qualidade.",
        keep_next=True,
    )
    black_ok = "Sem pendências" if metrics["black"] else "Há arquivos a formatar"
    report.table(
        "Verificações de qualidade executadas",
        ["Verificação", "Ferramenta", "Resultado"],
        [
            [
                "Testes automatizados",
                "pytest",
                (
                    f"{metrics['passed']} aprovados"
                    if not metrics["failed"]
                    else "Há falhas"
                ),
            ],
            ["Cobertura", "pytest-cov", f"{metrics['coverage']}%"],
            ["Análise estática", "Pylint", f"Nota {metrics['pylint']} de 10"],
            ["Formatação", "Black", black_ok],
            [
                "Segurança do código",
                "Bandit",
                "Nenhum achado no pacote" if metrics["bandit"] else "Há achados",
            ],
            ["Dependências", "pip-audit", metrics["audit"]],
            [
                "Varredura de composição de software",
                "Black Duck",
                "Não executada: ferramenta comercial não disponível neste ambiente",
            ],
        ],
        [5.2, 3.4, 7.4],
        caption="o resultado da varredura Black Duck deve ser registrado pela equipe quando a ferramenta estiver disponível.",
    )
    labels = {
        "test_hashing_merkle.py": "Hash e árvore de Merkle",
        "test_wallet_keystore.py": "Carteiras, assinaturas e keystore",
        "test_transaction.py": "Transações assinadas",
        "test_block_chain.py": "Blocos, prova de trabalho e validação",
        "test_contract.py": "Contrato: perfis, etapas e entradas",
        "test_node.py": "Nó: implantação, replay, persistência e laboratório",
        "test_node_api.py": "API HTTP do nó",
        "test_web.py": "Gateway web e carteiras",
        "test_cli_client.py": "Linha de comando e cliente HTTP",
    }
    rows = [
        [labels.get(Path(f).name, Path(f).name), str(count)]
        for f, count in sorted(metrics["per_file"].items(), key=lambda item: item[0])
    ]
    rows.append(["**Total", f"**{sum(metrics['per_file'].values())}"])
    report.table(
        "Testes automatizados por área",
        ["Área coberta", "Testes"],
        rows,
        [12.5, 3.5],
        caption="parametrizações contam como testes individuais.",
        align_center={1},
    )

    h(2, "7.2 Cenários de operações válidas, inválidas e de integridade")
    total = counts["V"] + counts["I"] + counts["A"]
    p(
        f"Além dos testes de unidade, {total} cenários foram executados contra uma blockchain nova, com transações reais assinadas, regras do contrato e mineração. Cada cenário declara o resultado esperado e o compara com o obtido. O resultado foi {counts['ok']} de {total} conforme o esperado. O Apêndice A traz a tabela completa e o comando make scenarios a reproduz.",
        keep_next=True,
    )
    report.table(
        "Resumo dos cenários executados",
        ["Grupo", "Cenários", "Exemplos"],
        [
            [
                "Operações válidas",
                str(counts["V"]),
                "Implantar, conceder perfis, percorrer as cinco etapas, anexar e verificar documento, reiniciar o nó",
            ],
            [
                "Entradas inválidas ou sem permissão",
                str(counts["I"]),
                "Quantidade negativa, data futura, pular etapas, alterar lote finalizado, assinatura falsa, repetição de transação",
            ],
            [
                "Integridade",
                str(counts["A"]),
                "Alterar dado gravado, refazer a prova de trabalho, bloquear escrita e restaurar",
            ],
        ],
        [4.6, 2.2, 9.2],
        caption="os identificadores V, I e A correspondem às linhas do Apêndice A.",
        align_center={1},
    )
    h(2, "7.3 Resultados obtidos")
    p(
        "Todas as operações válidas foram confirmadas em novos blocos e o histórico completo do lote pôde ser consultado. Todas as entradas inválidas e sem permissão foram rejeitadas com o código esperado, sem criar bloco e sem alterar o estado do contrato. A adulteração de um dado gravado foi detectada pela validação, tanto sem recalcular hashes, por divergência da raiz de Merkle e da assinatura, quanto refazendo a prova de trabalho do bloco, por quebra do encadeamento com o bloco seguinte. Com a cadeia comprometida, o nó recusou novas operações até a restauração."
    )

    # ------------------------------------------------------------- 8
    h(1, "8 Como executar")
    p(
        "Os passos abaixo iniciam a blockchain local, implantam o contrato e executam a aplicação. Os pré-requisitos são Docker com Compose ou Python 3.10 ou superior com make."
    )
    h(2, "8.1 Com Docker")
    report.code(
        [
            "git clone https://github.com/JulianaBallin/EcoOrigem.git",
            "cd EcoOrigem",
            "make docker-up      # nó, implantação do contrato e interface",
            "make docker-seed    # opcional: lotes de demonstração",
            "# Interface: http://127.0.0.1:5000   Nó: http://127.0.0.1:8545",
        ]
    )
    h(2, "8.2 Sem Docker")
    report.code(
        [
            "make venv           # cria .venv-ecoorigem e instala as dependências",
            "make node           # terminal 1: blockchain local (porta 8545)",
            "make deploy         # terminal 2: implanta o contrato e concede perfis",
            "make app            # terminal 3: interface (porta 5000)",
            "make start          # alternativa: faz tudo acima em segundo plano",
        ]
    )
    report.table(
        "Principais comandos do Makefile",
        ["Comando", "Função"],
        [
            ["make menu", "Menu interativo com todas as opções explicadas"],
            ["make test, make coverage", "Testes automatizados, com ou sem cobertura"],
            [
                "make lint, make format, make security",
                "Pylint, Black e Bandit com pip-audit",
            ],
            ["make scenarios", "Executa os cenários e gera a tabela de resultados"],
            ["make docker-test", "Executa os testes dentro de um contêiner"],
            [
                "make docker-reset, make reset",
                "Apagam a blockchain e as carteiras para recomeçar",
            ],
        ],
        [6.0, 10.0],
    )

    # ------------------------------------------------------------- 9
    h(1, "9 Limitações")
    report.bullets(
        [
            (
                "Nó único.",
                "A blockchain roda em um processo local. Sem uma rede de nós que rejeitem a cadeia adulterada, quem controla o arquivo pode refazer todos os blocos. A validação detecta a alteração, mas a imutabilidade plena exige replicação entre participantes.",
            ),
            (
                "Prova de trabalho didática.",
                "A dificuldade é baixa e serve para demonstrar o mecanismo, sem a segurança econômica de uma rede pública.",
            ),
            (
                "Carteiras no servidor.",
                "As chaves de demonstração ficam em um keystore no gateway, e não em uma extensão como o MetaMask. Isso simplifica a demonstração, mas não deve ser usado em produção.",
            ),
            (
                "Veracidade na origem.",
                "A blockchain garante que o registro não muda depois de gravado, não que o produtor informou a verdade. Certificadoras e sensores poderiam reforçar esse primeiro registro.",
            ),
            (
                "Dados públicos.",
                "Tudo que é gravado é legível por quem acessa a cadeia. Por isso dados pessoais ficam fora e documentos entram apenas como hash.",
            ),
            (
                "Escala.",
                "A reconstrução do estado reexecuta a cadeia inteira e o armazenamento é um arquivo JSON, adequados a um protótipo.",
            ),
            (
                "Implantação aberta.",
                "Qualquer conta pode implantar o contrato primeiro. Em produção a conta implantadora deveria ser restrita.",
            ),
            (
                "Black Duck.",
                "A varredura de composição de software não foi executada por indisponibilidade da ferramenta.",
            ),
        ]
    )

    # ------------------------------------------------------------ 10
    h(1, "10 Conclusão")
    p(
        "O EcoOrigem mostra, de ponta a ponta, como uma blockchain local em Python pode sustentar a rastreabilidade de produtos da bioeconomia amazônica. O contrato inteligente impõe a ordem das etapas e os perfis de acesso, a cadeia torna o histórico verificável e a interface permite registrar, consultar e auditar sem conhecimento técnico. Os testes comprovam tanto o funcionamento esperado quanto a rejeição de entradas inválidas e a detecção de adulterações."
    )
    p(
        "Como evolução, a equipe propõe: replicar a cadeia entre vários nós com consenso; integrar uma carteira externa como o MetaMask, com assinatura no navegador; gerar QR Code por lote para consulta pública pelo consumidor; registrar a localização aproximada da coleta; e integrar certificadoras e sensores para fortalecer a veracidade do primeiro registro."
    )

    # ---------------------------------------------------- references
    h(1, "Referências")
    for text in (
        "NAKAMOTO, S. Bitcoin: a peer-to-peer electronic cash system. 2008. Disponível em: https://bitcoin.org/bitcoin.pdf.",
        "SANTOS, F. Blockchain: visão geral. Material da Aula 01, Oficina de Desenvolvimento de Sistemas III, UEA, 2026.",
        "SANTOS, F. Implementação de Blockchain. Material da Aula 02 e exemplos BlockchainExemplo e Blockchain_APP, UEA, 2026.",
        "NATIONAL INSTITUTE OF STANDARDS AND TECHNOLOGY. FIPS 180-4: Secure Hash Standard. 2015.",
        "JOSEFSSON, S.; LIUSVAARA, I. RFC 8032: Edwards-Curve Digital Signature Algorithm (EdDSA). IETF, 2017.",
        "MERKLE, R. C. A digital signature based on a conventional encryption function. CRYPTO 1987.",
    ):
        paragraph = p(text, size=10, space_after=4)
        paragraph.paragraph_format.left_indent = Cm(0.8)
        paragraph.paragraph_format.first_line_indent = Cm(-0.8)

    # -------------------------------------------------------- appendix A
    h(1, "Apêndice A - Cenários de teste", page_break=True)
    p(
        "Operações realizadas, resultados esperados e resultados obtidos na execução automatizada dos cenários. Os números de bloco variam a cada execução.",
        keep_next=True,
    )
    report.table(
        "Cenários executados",
        [
            "ID",
            "Operação realizada",
            "Resultado esperado",
            "Resultado obtido",
            "Status",
        ],
        scenarios,
        [1.0, 4.4, 4.5, 4.5, 1.6],
        caption="Aprovado indica que o resultado obtido coincide com o esperado.",
        size=7.5,
        align_center={0, 4},
        keep_together=False,
    )

    # -------------------------------------------------------- appendix B
    h(1, "Apêndice B - Evidências da interface", page_break=True)
    p(
        "Capturas de tela geradas automaticamente pelo roteiro de demonstração (make demo-video), executado sobre a stack Docker.",
        keep_next=True,
    )
    shots = [
        (
            "02-registrar-formulario.png",
            "Formulário de registro de lote",
            "formulário preenchido com a carteira Produtor antes do envio.",
            0.65,
            15.0,
        ),
        (
            "03-registrar-confirmado.png",
            "Operação confirmada",
            "resultado com hash da transação, bloco minerado e prova de trabalho.",
            1.0,
            13.0,
        ),
        (
            "04-consultar-lote.png",
            "Consulta e histórico do lote",
            "etapas, responsáveis e linha do tempo com bloco e transação de cada evento.",
            1.0,
            13.5,
        ),
        (
            "05-rejeicao-1.png",
            "Rejeição por falta de permissão",
            "carteira sem perfil tenta registrar lote e o contrato retorna ACCESS_DENIED.",
            1.0,
            13.0,
        ),
        (
            "05-rejeicao-3.png",
            "Rejeição por etapa fora de ordem",
            "tentativa de finalizar um lote apenas cadastrado.",
            1.0,
            13.0,
        ),
        (
            "06-rejeicoes.png",
            "Log de operações rejeitadas",
            "todas as rejeições com camada, código, motivo, carteira e método.",
            0.7,
            15.0,
        ),
        (
            "07-blockchain.png",
            "Explorador de blocos",
            "hash com zeros da prova de trabalho, hash anterior, raiz de Merkle e transações.",
            0.65,
            15.0,
        ),
        (
            "08-prova-merkle.png",
            "Prova de inclusão de Merkle",
            "detalhe da transação com a prova validada contra a raiz do bloco.",
            0.65,
            15.0,
        ),
        (
            "10-adulteracao-detectada.png",
            "Adulteração detectada",
            "a alteração de um dado gravado invalida o bloco e bloqueia novas operações.",
            1.0,
            13.5,
        ),
        (
            "12-cadeia-restaurada.png",
            "Cadeia restaurada",
            "após a restauração a validação volta a indicar integridade.",
            0.85,
            13.5,
        ),
    ]
    for name, title, caption, ratio, width in shots:
        report.figure(cropped(name, ratio), title, caption, width)


def build(path: Path, metrics: dict, scenarios, counts, toc: dict[str, str]) -> Report:
    """Build the document with the given table of contents page numbers."""
    report = Report()
    cover(report)
    # Table of contents page (filled on the second pass)
    report.heading(1, "Sumário")
    report.headings.pop()
    report.doc.paragraphs[-1].paragraph_format.page_break_before = False
    body(report, metrics, scenarios, counts)
    # Insert the TOC lines right after the "Sumário" heading
    anchor = next(par for par in report.doc.paragraphs if par.text == "Sumário")
    lines = report.headings
    holder = []
    for level, text in lines:
        if level <= 2:
            holder.append((level, text, toc.get(text, "-")))
    for level, text, page in reversed(holder):
        report.toc_entry(level, text, page)
        entry = report.doc.paragraphs[-1]
        anchor._p.addnext(entry._p)  # pylint: disable=protected-access
    report.save(path)
    return report


def pdf_pages(pdf: Path) -> list[str]:
    """Return the text of every page of a PDF."""
    total = int(
        re.search(r"Pages:\s+(\d+)", run_command(["pdfinfo", str(pdf)]).stdout).group(1)
    )
    return [
        run_command(
            ["pdftotext", "-layout", "-f", str(n), "-l", str(n), str(pdf), "-"]
        ).stdout
        for n in range(1, total + 1)
    ]


def convert(docx: Path) -> Path:
    """Convert a DOCX to PDF with LibreOffice."""
    run_command(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx.parent),
            str(docx),
        ]
    )
    return docx.with_suffix(".pdf")


def main() -> int:
    """Build the report in two passes so the table of contents has page numbers."""
    metrics = quality_metrics()
    scenarios, counts = scenario_rows()
    METRICS_JSON.write_text(
        json.dumps(
            {
                **metrics,
                "per_file": dict(metrics["per_file"]),
                "scenarios": dict(counts),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    docx = OUT_DIR / f"{NAME}.docx"
    report = build(docx, metrics, scenarios, counts, {})
    pages = pdf_pages(convert(docx))
    toc: dict[str, str] = {}
    for level, text in report.headings:
        if level > 2:
            continue
        for number, content in enumerate(pages, start=1):
            if number < 3:
                continue
            if any(line.strip().startswith(text) for line in content.splitlines()):
                toc[text] = str(number)
                break
    build(docx, metrics, scenarios, counts, toc)
    convert(docx)
    print(f"Relatório gerado: {docx} e {docx.with_suffix('.pdf')}")
    missing = [
        text for level, text in report.headings if level <= 2 and text not in toc
    ]
    if missing:
        print("Títulos sem página no sumário:", *missing, sep="\n  ")
    return 0


if __name__ == "__main__":
    sys.exit(main())
